"""
Word 解析器测试文件，用于验证 Word 文档解析逻辑。
"""
import io

import pytest
from docx import Document as DocxDocument

from unittest.mock import MagicMock, patch, PropertyMock
from app.parsers.word_parser import WordParser


@pytest.mark.us25
def test_word_parser_extract_paragraphs():
    """
    测试 WordParser 是否能正确提取段落文本。
    """
    parser = WordParser()
    
    # 模拟 WordParser 模块中导入的 Document
    with patch("app.parsers.word_parser.Document") as mock_document:
        mock_doc = MagicMock()
        mock_document.return_value = mock_doc
        
        # 模拟段落
        mock_para_1 = MagicMock()
        mock_para_1.text = "这是第一段文本"
        mock_para_1.style.name = "Normal"
        
        mock_para_2 = MagicMock()
        mock_para_2.text = "这是第二段文本"
        mock_para_2.style.name = "Normal"
        
        mock_doc.paragraphs = [mock_para_1, mock_para_2]
        mock_doc.tables = []
        
        # 执行解析
        result = parser.parse(b"fake docx binary stream")
        
        # 断言校验
        assert len(result) == 1
        assert result[0]["page_number"] == 1
        assert "第一段文本" in result[0]["content"]
        assert "第二段文本" in result[0]["content"]
        assert result[0]["parser_name"] == "WordParser"
        assert result[0]["parser_version"] == "1.0.0"
        assert result[0]["metadata"]["source_type"] == "word"
        assert result[0]["metadata"]["paragraphs_count"] == 2
        assert result[0]["metadata"]["tables_count"] == 0


@pytest.mark.us25
def test_word_parser_extract_tables():
    """
    测试 WordParser 是否能正确提取表格文本。
    """
    parser = WordParser()
    
    with patch("app.parsers.word_parser.Document") as mock_document:
        mock_doc = MagicMock()
        mock_document.return_value = mock_doc
        
        mock_doc.paragraphs = []
        
        # 模拟表格
        mock_cell_1 = MagicMock()
        mock_cell_1.text = "姓名"
        mock_cell_2 = MagicMock()
        mock_cell_2.text = "年龄"
        mock_cell_3 = MagicMock()
        mock_cell_3.text = "张三"
        mock_cell_4 = MagicMock()
        mock_cell_4.text = "25"
        
        mock_row_1 = MagicMock()
        mock_row_1.cells = [mock_cell_1, mock_cell_2]
        mock_row_2 = MagicMock()
        mock_row_2.cells = [mock_cell_3, mock_cell_4]
        
        mock_table = MagicMock()
        mock_table.rows = [mock_row_1, mock_row_2]
        mock_doc.tables = [mock_table]
        
        # 执行解析
        result = parser.parse(b"fake docx binary stream")
        
        # 断言校验
        assert len(result) == 1
        assert "姓名" in result[0]["content"]
        assert "年龄" in result[0]["content"]
        assert "张三" in result[0]["content"]
        assert "25" in result[0]["content"]
        assert result[0]["parser_name"] == "WordParser"
        assert result[0]["parser_version"] == "1.0.0"
        assert result[0]["metadata"]["tables_count"] == 1


@pytest.mark.us25
def test_word_parser_headings():
    """
    测试 WordParser 是否能正确识别标题层级。
    """
    parser = WordParser()
    
    with patch("app.parsers.word_parser.Document") as mock_document:
        mock_doc = MagicMock()
        mock_document.return_value = mock_doc
        
        # 模拟带标题的段落
        mock_heading_1 = MagicMock()
        mock_heading_1.text = "第一章 引言"
        mock_heading_1.style.name = "Heading 1"
        
        mock_heading_2 = MagicMock()
        mock_heading_2.text = "1.1 研究背景"
        mock_heading_2.style.name = "Heading 2"
        
        mock_para = MagicMock()
        mock_para.text = "这是一些正文内容"
        mock_para.style.name = "Normal"
        
        mock_doc.paragraphs = [mock_heading_1, mock_heading_2, mock_para]
        mock_doc.tables = []
        
        # 执行解析
        result = parser.parse(b"fake docx binary stream")
        
        # 断言校验
        headings = result[0]["metadata"]["headings"]
        assert len(headings) == 2
        assert headings[0]["level"] == 1
        assert headings[0]["text"] == "第一章 引言"
        assert headings[1]["level"] == 2
        assert headings[1]["text"] == "1.1 研究背景"
        assert result[0]["parser_name"] == "WordParser"
        assert result[0]["parser_version"] == "1.0.0"


@pytest.mark.us25
def test_word_parser_empty_document():
    """
    测试 WordParser 处理空文档的边界情况。
    """
    parser = WordParser()
    
    with patch("app.parsers.word_parser.Document") as mock_document:
        mock_doc = MagicMock()
        mock_document.return_value = mock_doc
        
        mock_doc.paragraphs = []
        mock_doc.tables = []
        
        # 执行解析
        result = parser.parse(b"fake docx binary stream")
        
        # 断言校验
        assert len(result) == 1
        assert result[0]["content"] == ""
        assert result[0]["char_count"] == 0
        assert result[0]["parser_name"] == "WordParser"
        assert result[0]["parser_version"] == "1.0.0"
        assert result[0]["metadata"]["paragraphs_count"] == 0
        assert result[0]["metadata"]["tables_count"] == 0


@pytest.mark.us25
def test_word_parser_skip_empty_paragraph():
    """
    测试 WordParser 跳过空文本段落（覆盖第 32 行 continue）。
    """
    parser = WordParser()
    
    with patch("app.parsers.word_parser.Document") as mock_document:
        mock_doc = MagicMock()
        mock_document.return_value = mock_doc
        
        # 模拟包含空文本的段落
        mock_para_empty = MagicMock()
        mock_para_empty.text = "   "  # 空白文本，应被跳过
        mock_para_empty.style.name = "Normal"
        
        mock_para_valid = MagicMock()
        mock_para_valid.text = "有效内容"
        mock_para_valid.style.name = "Normal"
        
        mock_doc.paragraphs = [mock_para_empty, mock_para_valid]
        mock_doc.tables = []
        
        # 执行解析
        result = parser.parse(b"fake docx binary stream")
        
        # 断言校验
        assert len(result) == 1
        assert "有效内容" in result[0]["content"]
        assert result[0]["metadata"]["paragraphs_count"] == 1


@pytest.mark.us25
def test_word_parser_reads_real_document_bytes():
    """
    用**真实** .docx 字节解析，段落与表格都要被读到。

    上面几个用例把 Document 整个 mock 掉了，验证的是 WordParser 自己的分支逻辑；
    这个用例补上另一端：真实 OOXML 字节进来，内容确实能被提取。
    此前仓库里唯一的 Word 文件是一个 0 字节的 empty_file.docx，且没有任何测试引用它。
    """
    buffer = io.BytesIO()
    document = DocxDocument()
    document.add_heading("Compte rendu de travaux pratiques", level=1)
    document.add_paragraph("Premier paragraphe du corps du texte.")
    document.add_paragraph("Second paragraphe du corps du texte.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "modele"
    table.cell(0, 1).text = "score"
    table.cell(1, 0).text = "foret"
    table.cell(1, 1).text = "0.84"
    document.save(buffer)

    result = WordParser().parse(buffer.getvalue())

    assert len(result) == 1
    unit = result[0]
    assert unit["parser_name"] == "WordParser"
    assert "Compte rendu de travaux pratiques" in unit["content"]
    assert "Premier paragraphe du corps du texte." in unit["content"]
    assert "Second paragraphe du corps du texte." in unit["content"]
    assert "foret | 0.84" in unit["content"]  # 表格行按 " | " 拼接
    assert unit["metadata"]["tables_count"] == 1
    assert unit["metadata"]["paragraphs_count"] >= 2
    assert unit["char_count"] == len(unit["content"])